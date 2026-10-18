import json

import docx
from django.core.management import BaseCommand

from home.utils import docx_parse_document
from home.utils import eaf_to_docx

EAF_FILE_NAME = 'Ani.eaf'
FILE_NAME_WITHOUT_EXTENSION = EAF_FILE_NAME.rsplit('.', 1)[0]


class Command(BaseCommand):
    def add_arguments(self, parser):
        """
        :type parser:  argparse.ArgumentParser
        """
        parser.add_argument('-f', '--filepath', dest='filepath', help='Path to .eaf, .docx, or .ass file')
        parser.add_argument('-s', '--separate', dest='separate_speech', default=True, action='store_true')

    def handle(self, *args, **options):
        main(file_path=options['filepath'], separate_speech=options['separate_speech'])


def main(file_path, separate_speech=True):
    file_name_without_extension = file_path.rsplit('.', 1)[0]
    if file_path.lower().endswith('.docx'):
        with open(file_path, 'rb') as f:
            document = docx.Document(f)
    elif file_path.lower().endswith('.ass'):
        from home.utils import ass_to_docx
        with open(file_path, 'r', encoding='utf-8') as ass_file:
            docx_file_obj = ass_to_docx(ass_file, separate_speech)
        with open(f'{file_name_without_extension}.docx', 'wb') as docx_file:
            docx_file.write(docx_file_obj.read())
        with open(f'{file_name_without_extension}.docx', 'rb') as f:
            document = docx.Document(f)
    else:
        with open(file_path, 'rb') as eaf_file:
            eaf_docx = eaf_to_docx(eaf_file, True, separate_speech)
        with open(f'{file_name_without_extension}.docx',
                  'wb') as docx_file:
            docx_file.write(eaf_docx.read())

        with open(f'{file_name_without_extension}.docx', 'rb') as f:
            document = docx.Document(f)

    html, transcript_map = docx_parse_document(document)

    # Full master JSON with all metadata
    with open(f'{file_name_without_extension}.json', 'w', encoding='utf-8') as map_json_file:
        json.dump(transcript_map, map_json_file, indent=2, ensure_ascii=False)

    # Minified compact JSON for Wagtail admin (id, start, end, speaker)
    admin_map = [
        {
            'id': item['id'],
            'start': int(item['start']),
            'end': int(item['end']),
            'speaker': item.get('speaker', '')
        }
        for item in transcript_map
    ]
    with open(f'{file_name_without_extension}-admin.json', 'w', encoding='utf-8') as admin_json_file:
        json.dump(admin_map, admin_json_file, ensure_ascii=False, separators=(',', ':'))

    # Format HTML with clean newlines after block tags
    formatted_html = html
    for tag in ['</p>', '</h3>', '</h5>', '</div>', '</blockquote>']:
        formatted_html = formatted_html.replace(tag, tag + '\n')

    with open(f'{file_name_without_extension}.html', 'w', encoding='utf-8') as html_file:
        html_file.write(formatted_html)

    admin_map_json_str = json.dumps(admin_map, ensure_ascii=False)
    preview_html = (
        '<!DOCTYPE html>\n'
        '<html lang="fr">\n'
        '<head>\n'
        '    <meta charset="UTF-8">\n'
        '    <meta name="viewport" content="width=device-width, initial-scale=1.0">\n'
        '    <title>Preview: Transcript & Player</title>\n'
        '    <link rel="stylesheet" href="https://cdn.jsdelivr.net/npm/bootstrap@4.6.2/dist/css/bootstrap.min.css">\n'
        '    <style>\n'
        '        body { padding: 20px 20px 100px 20px; max-width: 900px; margin: 0 auto; font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, "Helvetica Neue", Arial, sans-serif; line-height: 1.6; color: #212529; }\n'
        '        .sticky-player { position: sticky; top: 0; background: rgba(255, 255, 255, 0.95); backdrop-filter: blur(5px); padding: 12px 15px; border-bottom: 2px solid #dee2e6; z-index: 1000; margin-bottom: 25px; box-shadow: 0 4px 6px -1px rgba(0,0,0,0.1); border-radius: 6px; }\n'
        '        .sticky-player audio { width: 100%; height: 40px; }\n'
        '        .bs-callout { padding: 15px; border: 1px solid #eee; border-left-width: 5px; border-radius: 4px; margin: 15px 0; }\n'
        '        .bs-callout-info { border-left-color: #17a2b8; background-color: #f1f3f5; font-style: italic; font-weight: 500; margin-left: 20px; }\n'
        '        blockquote.font-italic { margin-top: 1rem; margin-bottom: 1rem; padding-left: 1.5rem; border-left: 4px solid #ced4da; }\n'
        '        .transcript-line { transition: background-color 0.15s, color 0.15s; border-radius: 3px; padding: 1px 3px; }\n'
        '        .transcript-line:hover { background-color: #fff9c4 !important; color: #212529 !important; cursor: pointer; text-decoration: underline; }\n'
        '        .bg-primary { background-color: #007bff !important; }\n'
        '        .text-white { color: #fff !important; }\n'
        '    </style>\n'
        '</head>\n'
        '<body>\n'
        '    <div class="sticky-player">\n'
        '        <div class="d-flex justify-content-between align-items-center mb-1">\n'
        '            <strong>Аудиотранскрипт (' + str(len(admin_map)) + ' реплик)</strong>\n'
        '            <span class="badge badge-info" id="timeDisplay">00:00 / 00:00</span>\n'
        '        </div>\n'
        '        <audio id="newAudioTranscript" controls preload="metadata">\n'
        '            <source src="https://files.le-francais.ru/lecons/restored/french-009-restored.mp3" type="audio/mpeg">\n'
        '        </audio>\n'
        '        <small class="text-muted d-block mt-1">Клик по реплике перематывает плеер. При воспроизведении реплики подсвечиваются синим.</small>\n'
        '    </div>\n'
        '    <div class="content">\n'
        + formatted_html + '\n'
        '    </div>\n'
        '    <script src="https://code.jquery.com/jquery-3.6.0.min.js"></script>\n'
        '    <script>\n'
        '        let audio = $(\'#newAudioTranscript\');\n'
        '        let startsEndsMap = ' + admin_map_json_str + ';\n'
        '        audio.on(\'timeupdate\', function () {\n'
        '            let current = this.currentTime * 1000;\n'
        '            startsEndsMap.filter((item) => item.start <= current && item.end >= current).forEach((item) => {\n'
        '                $(`#${item.id}`).addClass([\'bg-primary\', \'text-white\']);\n'
        '            });\n'
        '            startsEndsMap.filter((item) => item.end <= current || item.start >= current).forEach((item) => {\n'
        '                $(`#${item.id}`).removeClass([\'bg-primary\', \'text-white\']);\n'
        '            });\n'
        '            let curSec = Math.floor(this.currentTime);\n'
        '            let durSec = Math.floor(this.duration || 0);\n'
        '            let fmt = (s) => String(Math.floor(s/60)).padStart(2, \'0\') + \':\' + String(s%60).padStart(2, \'0\');\n'
        '            $(\'#timeDisplay\').text(fmt(curSec) + \' / \' + fmt(durSec));\n'
        '        });\n'
        '        $(\'.transcript-line\').on(\'click\', function (){\n'
        '            let start = $(this).data(\'start\');\n'
        '            if (start !== undefined) {\n'
        '                audio[0].currentTime = start / 1000;\n'
        '                audio[0].play();\n'
        '            }\n'
        '        });\n'
        '    </script>\n'
        '</body>\n'
        '</html>\n'
    )
    with open(f'{file_name_without_extension}_preview.html', 'w', encoding='utf-8') as preview_file:
        preview_file.write(preview_html)


