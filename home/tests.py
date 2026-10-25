from django.test import TestCase
from home.models import PageWithSidebar, ArticlePage

class Page2ArticleTestCase(TestCase):
    def setUp(self):
        PageWithSidebar.objects.create(
            path='00010001',
            depth='2',
            slug='slug',
            url_path='home',
            title='title',
            seo_title='seo_title',
            show_in_menus=True,
            menu_title='menu_title',
            is_nav_root=False,
            is_selectable=True,
            reference_title='reference_title',
            subtitle='subtitle',
            live=True,
        )


import io
import os

import docx
from lxml import etree
from home.utils import docx_parse_document


class DocxParserTestCase(TestCase):
    def test_paragraph_styles_and_block_transitions(self):
        doc = docx.Document()
        doc.add_paragraph('Heading text', style='Heading 1')
        doc.add_paragraph('Quote 1', style='Quote')
        doc.add_paragraph('Quote 2', style='Quote')
        doc.add_paragraph('Normal paragraph')
        doc.add_paragraph('Callout 1', style='Intense Quote')
        doc.add_paragraph('Callout 2', style='Intense Quote')
        doc.add_paragraph('List item 1', style='List Paragraph')
        doc.add_paragraph('List item 2', style='List Paragraph')
        doc.add_paragraph('Final paragraph')

        html, lines_map = docx_parse_document(doc)
        self.assertIn('<h2>Heading text</h2>', html)
        self.assertIn('<blockquote class="font-italic pl-4"><p>Quote 1</p><p>Quote 2</p></blockquote>', html)
        self.assertIn('<p>Normal paragraph</p>', html)
        self.assertIn('<div class="bs-callout bs-callout-info bs-callout-transcript"><p>Callout 1</p><p>Callout 2</p></div>', html)
        self.assertIn('<ul><li>List item 1</li><li>List item 2</li></ul>', html)
        self.assertIn('<p>Final paragraph</p>', html)

    def test_goback_bookmark_does_not_leak_or_break_opened_tags(self):
        doc = docx.Document()
        p = doc.add_paragraph()
        run = p.add_run('Text with GoBack')

        b_start = etree.Element('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}bookmarkStart',
                                {'{http://schemas.openxmlformats.org/wordprocessingml/2006/main}id': '0',
                                 '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}name': '_GoBack'})
        b_end = etree.Element('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}bookmarkEnd',
                              {'{http://schemas.openxmlformats.org/wordprocessingml/2006/main}id': '0'})
        p._p.insert(0, b_start)
        p._p.append(b_end)

        html, lines_map = docx_parse_document(doc)
        self.assertNotIn('_GoBack', html)
        self.assertNotIn('<span id="_GoBack">', html)

    def test_bookmark_with_name_creates_span(self):
        doc = docx.Document()
        p = doc.add_paragraph()
        run = p.add_run('Bookmarked text')

        b_start = etree.Element('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}bookmarkStart',
                                {'{http://schemas.openxmlformats.org/wordprocessingml/2006/main}id': '1',
                                 '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}name': 'my_anchor'})
        b_end = etree.Element('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}bookmarkEnd',
                              {'{http://schemas.openxmlformats.org/wordprocessingml/2006/main}id': '1'})
        p._p.insert(0, b_start)
        p._p.append(b_end)

        html, lines_map = docx_parse_document(doc)
        self.assertIn('<span id="my_anchor">', html)
        self.assertIn('</span>', html)

    def test_non_timing_comment_is_balanced_and_not_leaked(self):
        doc = docx.Document()
        p = doc.add_paragraph('Commented text')
        p.add_comment('Fix&"later" <soon>', author='editor')

        html, lines_map = docx_parse_document(doc)
        self.assertNotIn('<span class="comment">', html)
        self.assertEqual(html.count('<span'), html.count('</span>'))
        self.assertIn('class="arnaud-comment"', html)
        self.assertIn('data-content="Fix&amp;&quot;later&quot; &lt;soon&gt;"', html)
        self.assertIn('data-toggle="popover"', html)
        self.assertIn('data-container="body"', html)
        self.assertIn('data-placement="top"', html)
        self.assertEqual(lines_map, [])
        self.assertEqual(len(lines_map.comments), 1)
        self.assertEqual(lines_map.comments[0]['text'], 'Fix&"later" <soon>')
        self.assertEqual(lines_map.comments[0]['author'], 'editor')

    def test_non_timing_comment_inside_replica_does_not_break_tags_and_is_captured(self):
        doc = docx.Document()
        p = doc.add_paragraph('— Bonjour tout le monde !')
        p.runs[0].add_comment('Teacher 0 1000')
        p.add_comment('Note for editor', author='editor')

        html, lines_map = docx_parse_document(doc)
        self.assertNotIn('<span class="comment">', html)
        self.assertEqual(html.count('<span'), html.count('</span>'))
        self.assertIn('class="arnaud-comment"', html)
        self.assertIn('data-content="Note for editor"', html)
        self.assertEqual(len(lines_map), 1)
        self.assertIn('comments', lines_map[0])
        self.assertEqual(lines_map[0]['comments'][0]['text'], 'Note for editor')
        self.assertEqual(len(lines_map.comments), 1)
        self.assertEqual(lines_map.comments[0]['text'], 'Note for editor')

    def test_interjection_with_dialogue_dash_is_attributed_to_unknown(self):
        doc = docx.Document()
        p1 = doc.add_paragraph('— Première partie du cours.')
        p1.runs[0].add_comment('Teacher 0 100')

        p2 = doc.add_paragraph('— Oui, madame ?')
        p2.runs[0].add_comment('Teacher 100 200')

        p3 = doc.add_paragraph('— Deuxième partie du cours.')
        p3.runs[0].add_comment('Teacher 200 300')

        html, lines_map = docx_parse_document(doc)
        self.assertEqual(len(lines_map), 3)
        self.assertEqual(lines_map[0]['speaker'], 'Teacher')
        self.assertEqual(lines_map[1]['speaker'], 'unknown')
        self.assertEqual(lines_map[2]['speaker'], 'Teacher')

        self.assertIn('data-speaker="Teacher"', html)
        self.assertIn('data-speaker="unknown"', html)

    def test_normal_dialogue_alternation_does_not_become_unknown(self):
        doc = docx.Document()
        p1 = doc.add_paragraph('— Bonjour.')
        p1.runs[0].add_comment('Teacher 0 100')

        p2 = doc.add_paragraph('— Bonjour, madame.')
        p2.runs[0].add_comment('Student 100 200')

        p3 = doc.add_paragraph('— Asseyez-vous.')
        p3.runs[0].add_comment('Teacher 200 300')

        html, lines_map = docx_parse_document(doc)
        self.assertEqual(len(lines_map), 3)
        self.assertEqual(lines_map[0]['speaker'], 'Teacher')
        self.assertEqual(lines_map[1]['speaker'], 'Student')
        self.assertEqual(lines_map[2]['speaker'], 'Teacher')
        self.assertNotIn('unknown', html)

    def test_continuation_without_dash_after_unknown_remains_unknown(self):
        doc = docx.Document()
        p1 = doc.add_paragraph('— Début du professeur.')
        p1.runs[0].add_comment('Teacher 0 100')

        p2 = doc.add_paragraph('— Interjection élève.')
        p2.runs[0].add_comment('Teacher 100 200')

        p3 = doc.add_paragraph('Suite de l\'élève sans tiret.')
        p3.runs[0].add_comment('Teacher 200 300')

        p4 = doc.add_paragraph('— Reprise du professeur avec tiret.')
        p4.runs[0].add_comment('Teacher 300 400')

        html, lines_map = docx_parse_document(doc)
        self.assertEqual(len(lines_map), 4)
        self.assertEqual(lines_map[0]['speaker'], 'Teacher')
        self.assertEqual(lines_map[1]['speaker'], 'unknown')
        self.assertEqual(lines_map[2]['speaker'], 'unknown')
        self.assertEqual(lines_map[3]['speaker'], 'Teacher')

    def test_teacher_speech_continuation_without_dash_remains_teacher(self):
        doc = docx.Document()
        p1 = doc.add_paragraph('— Début du cours magistral.')
        p1.runs[0].add_comment('Teacher 0 100')

        p2 = doc.add_paragraph('Deuxième paragraphe sans tiret.')
        p2.runs[0].add_comment('Teacher 100 200')

        p3 = doc.add_paragraph('Troisième paragraphe sans tiret.')
        p3.runs[0].add_comment('Teacher 200 300')

        html, lines_map = docx_parse_document(doc)
        self.assertEqual(len(lines_map), 3)
        self.assertEqual(lines_map[0]['speaker'], 'Teacher')
        self.assertEqual(lines_map[1]['speaker'], 'Teacher')
        self.assertEqual(lines_map[2]['speaker'], 'Teacher')
        self.assertNotIn('unknown', html)

    def test_bookmarks_parsed_into_transcript_map_and_reproduced_in_docx(self):
        from home.utils import map_to_docx
        doc = docx.Document()
        p1 = doc.add_paragraph('— Première réplique avec signet.')
        r1 = p1.runs[0]
        r1.add_comment('Teacher 0 100')

        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        b_start = OxmlElement('w:bookmarkStart')
        b_start.set(qn('w:id'), '10')
        b_start.set(qn('w:name'), 'chapitre_1')
        p1._p.insert(0, b_start)
        b_end = OxmlElement('w:bookmarkEnd')
        b_end.set(qn('w:id'), '10')
        p1._p.append(b_end)

        html, lines_map = docx_parse_document(doc)
        self.assertEqual(len(lines_map), 1)
        self.assertEqual(lines_map[0]['bookmark'], 'chapitre_1')
        self.assertIn('chapitre_1', lines_map[0]['bookmarks'])
        self.assertEqual(lines_map[0]['text'], '— Première réplique avec signet.')
        self.assertIn('<span id="chapitre_1">', html)

        docx_file = map_to_docx(lines_map)
        roundtrip_doc = docx.Document(docx_file)
        roundtrip_html, roundtrip_map = docx_parse_document(roundtrip_doc)

        self.assertEqual(len(roundtrip_map), 1)
        self.assertEqual(roundtrip_map[0]['bookmark'], 'chapitre_1')
        self.assertIn('chapitre_1', roundtrip_map[0]['bookmarks'])
        self.assertEqual(roundtrip_map[0]['speaker'], 'Teacher')
        self.assertIn('<span id="chapitre_1">', roundtrip_html)

    def test_interjection_with_multiple_runs_in_paragraph(self):
        doc = docx.Document()
        p1 = doc.add_paragraph('— Discours prof.')
        p1.runs[0].add_comment('Teacher 0 100')

        p2 = doc.add_paragraph('— Réplique élève phrase 1. ')
        p2.runs[0].add_comment('Teacher 100 200')
        r2 = p2.add_run('Réplique élève phrase 2.')
        r2.add_comment('Teacher 200 300')

        p3 = doc.add_paragraph('— Suite prof.')
        p3.runs[0].add_comment('Teacher 300 400')

        html, lines_map = docx_parse_document(doc)
        self.assertEqual(len(lines_map), 4)
        self.assertEqual(lines_map[0]['speaker'], 'Teacher')
        self.assertEqual(lines_map[1]['speaker'], 'unknown')
        self.assertEqual(lines_map[2]['speaker'], 'unknown')
        self.assertEqual(lines_map[3]['speaker'], 'Teacher')

    def test_pipe_marker_splits_replica_proportionally(self):
        doc = docx.Document()
        p1 = doc.add_paragraph('— Première phrase courte. || Deuxième phrase un peu plus longue.')
        p1.runs[0].add_comment('Teacher 0 1000')

        html, lines_map = docx_parse_document(doc)
        self.assertEqual(len(lines_map), 2)
        self.assertEqual(lines_map[0]['speaker'], 'Teacher')
        self.assertEqual(lines_map[0]['text'], '— Première phrase courte.')
        self.assertEqual(lines_map[0]['start'], '0')

        self.assertEqual(lines_map[1]['speaker'], 'Teacher')
        self.assertEqual(lines_map[1]['text'], 'Deuxième phrase un peu plus longue.')
        self.assertEqual(lines_map[1]['end'], '1000')
        self.assertTrue(lines_map[1]['is_split'])
        self.assertTrue(lines_map[1]['split_adjacent'])

        # End of part 1 equals start of part 2
        self.assertEqual(lines_map[0]['end'], lines_map[1]['start'])
        self.assertNotIn('||', html)

        # Verify that HTML DOM was synchronized with split spans
        self.assertIn(f'id="{lines_map[0]["id"]}"', html)
        self.assertIn(f'id="{lines_map[1]["id"]}"', html)
        self.assertIn(f'data-end="{lines_map[0]["end"]}"', html)
        self.assertIn(f'data-start="{lines_map[1]["start"]}"', html)

    def test_ass_to_map_and_docx_roundtrip(self):
        from home.utils import ass_to_map, map_to_docx

        ass_content = (
            "[Script Info]\n"
            "Title: Test\n"
            "[V4+ Styles]\n"
            "Format: Name, Fontname, Fontsize, PrimaryColour, SecondaryColour, OutlineColour, BackColour, Bold, Italic, Underline, StrikeOut, ScaleX, ScaleY, Spacing, Angle, BorderStyle, Outline, Shadow, Alignment, MarginL, MarginR, MarginV, Encoding\n"
            "Style: Default,Arial,42,&H00FFFFFF,&H000000FF,&H00000000,&H00000000,0,0,0,0,100,100,0,0,1,2,2,2,10,10,10,1\n"
            "Style: Наташа,Arial,42,&H00FFFFFF,&H000000FF,&H00000000,&H00000000,0,0,0,0,100,100,0,0,1,2,2,2,10,10,10,1\n"
            "[Events]\n"
            "Format: Layer, Start, End, Style, Name, MarginL, MarginR, MarginV, Effect, Text\n"
            "Comment: 0,0:00:00.05,0:00:03.04,Bookmark,,0,0,0,bookmark,=== BOOKMARK: Vocabulaire ===\n"
            "Dialogue: 0,0:00:00.05,0:00:03.04,Наташа,Наташа,0,0,0,BOOKMARK:Vocabulaire,Bonjour à tous !\n"
            "Dialogue: 0,0:00:03.05,0:00:05.10,unknown,Миша,0,0,0,SPLIT_1,{\\c&H0000FFFF&}Bonjour professeur.\n"
        )
        t_map = ass_to_map(ass_content)
        self.assertEqual(len(t_map), 2)
        # Entry 1
        self.assertEqual(t_map[0]['speaker'], 'Наташа')
        self.assertEqual(t_map[0]['start'], '50')
        self.assertEqual(t_map[0]['end'], '3040')
        self.assertEqual(t_map[0]['text'], 'Bonjour à tous\u202f!')
        self.assertEqual(t_map[0]['bookmark'], 'Vocabulaire')

        # Entry 2 (edited in SE: actor changed to Миша, visual tag stripped)
        self.assertEqual(t_map[1]['speaker'], 'Миша')
        self.assertEqual(t_map[1]['start'], '3050')
        self.assertEqual(t_map[1]['end'], '5100')
        self.assertEqual(t_map[1]['text'], 'Bonjour professeur.')

        # Roundtrip to DOCX
        docx_bytes = map_to_docx(t_map)
        doc = docx.Document(docx_bytes)
        html, parsed_map = docx_parse_document(doc)

        self.assertEqual(len(parsed_map), 2)
        self.assertEqual(parsed_map[0]['speaker'], 'Наташа')
        self.assertEqual(parsed_map[0]['text'], 'Bonjour à tous\u202f!')
        self.assertEqual(parsed_map[0]['bookmark'], 'Vocabulaire')

        self.assertEqual(parsed_map[1]['speaker'], 'Миша')
        self.assertEqual(parsed_map[1]['text'], 'Bonjour professeur.')

    def test_sticky_material_hyperlink_unwrapped(self):
        from docx.oxml import parse_xml
        from docx.oxml.ns import nsdecls
        doc = docx.Document()

        # Add a Callout paragraph with an internal hyperlink
        p_callout = doc.add_paragraph(style='Intense Quote')
        p_callout.add_run('Attention : ')
        hl_xml = parse_xml(
            f'<w:hyperlink {nsdecls("w")} w:anchor="anchor_point_1">'
            f'  <w:r><w:rPr><w:b/></w:rPr><w:t>Règle importante</w:t></w:r>'
            f'</w:hyperlink>'
        )
        p_callout._p.append(hl_xml)

        # Downstream paragraph with the target bookmark
        p_target = doc.add_paragraph('Suite du cours...')
        b_start = parse_xml(f'<w:bookmarkStart {nsdecls("w")} w:id="101" w:name="anchor_point_1"/>')
        b_end = parse_xml(f'<w:bookmarkEnd {nsdecls("w")} w:id="101"/>')
        p_target._p.insert(0, b_start)
        p_target._p.append(b_end)

        html, lines_map = docx_parse_document(doc)

        # 1. The callout container must have data-sticky-until="anchor_point_1"
        self.assertIn('data-sticky-until="anchor_point_1"', html)
        self.assertIn('<div class="bs-callout bs-callout-info bs-callout-transcript"', html)

        # 2. Hyperlink must be unwrapped (no <a> or <w:hyperlink> tag in HTML)
        self.assertNotIn('<a ', html)
        self.assertNotIn('</a>', html)

        # 3. Formatting of text inside the hyperlink must be preserved
        self.assertIn('<strong>Règle importante</strong>', html)
        self.assertIn('Attention : ', html)

    def test_sticky_anchor_missing_raises_error(self):
        from docx.oxml import parse_xml
        from docx.oxml.ns import nsdecls
        doc = docx.Document()

        # Callout paragraph references anchor "non_existent_bookmark" which doesn't exist anywhere
        p_callout = doc.add_paragraph(style='Intense Quote')
        hl_xml = parse_xml(
            f'<w:hyperlink {nsdecls("w")} w:anchor="non_existent_bookmark">'
            f'  <w:r><w:t>Règle avec signet manquant</w:t></w:r>'
            f'</w:hyperlink>'
        )
        p_callout._p.append(hl_xml)

        with self.assertRaises(ValueError) as ctx:
            docx_parse_document(doc, strict=True)

        self.assertIn('non_existent_bookmark', str(ctx.exception))

    def test_sticky_anchor_missing_is_dropped_when_not_strict(self):
        doc = docx.Document()
        p_callout = doc.add_paragraph(style='Intense Quote')
        p_callout._p.append(self._hyperlink('non_existent_bookmark', 'Règle avec signet manquant'))

        with self.assertLogs('home.utils', level='WARNING') as cm:
            html, lines_map = docx_parse_document(doc)

        self.assertTrue(any('non_existent_bookmark' in msg for msg in cm.output))
        self.assertNotIn('data-sticky-until', html)
        self.assertIn('Règle avec signet manquant', html)

    @staticmethod
    def _hyperlink(anchor, text):
        from docx.oxml import parse_xml
        from docx.oxml.ns import nsdecls
        return parse_xml(
            f'<w:hyperlink {nsdecls("w")} w:anchor="{anchor}"><w:r><w:t>{text}</w:t></w:r></w:hyperlink>'
        )

    @staticmethod
    def _comment_runs(paragraph, text, first_run, last_run, author='Teacher'):
        """Add a Word comment whose range spans runs[first_run..last_run] (ranges may overlap)."""
        from docx.oxml.ns import qn
        runs = [r for r in paragraph.runs if r.text]
        comment = runs[first_run].add_comment(text, author=author)
        for end in paragraph._p.iter(qn('w:commentRangeEnd')):
            if end.get(qn('w:id')) == str(comment._id):
                runs[last_run]._r.addnext(end)
                break

    @staticmethod
    def _add_bookmark(paragraph, name, b_id):
        from docx.oxml import parse_xml
        from docx.oxml.ns import nsdecls
        paragraph._p.insert(0, parse_xml(f'<w:bookmarkStart {nsdecls("w")} w:id="{b_id}" w:name="{name}"/>'))
        paragraph._p.append(parse_xml(f'<w:bookmarkEnd {nsdecls("w")} w:id="{b_id}"/>'))

    def test_link_in_regular_paragraph_does_not_leak_into_next_callout(self):
        doc = docx.Document()
        p_link = doc.add_paragraph('Voir ')
        p_link._p.append(self._hyperlink('chapitre2', 'le chapitre 2'))
        doc.add_paragraph('Paragraphe intermédiaire.')
        doc.add_paragraph('Encadré sans lien.', style='Intense Quote')
        self._add_bookmark(doc.add_paragraph('Chapitre 2'), 'chapitre2', 301)

        html, lines_map = docx_parse_document(doc)

        self.assertIn('<div class="bs-callout bs-callout-info bs-callout-transcript"><p>Encadré sans lien.', html)
        # Regular sticky paragraph is wrapped in a div around (not inside) the <p>
        self.assertIn('<div data-sticky-until="chapitre2"><p>Voir le chapitre 2</p></div>', html)
        self.assertNotIn('<p><div', html)

    def test_hidden_word_bookmarks_are_not_sticky(self):
        doc = docx.Document()
        p_toc = doc.add_paragraph()
        p_toc._p.append(self._hyperlink('_Toc12345', 'Chapitre 1'))
        self._add_bookmark(doc.add_paragraph('Chapitre 1'), '_Toc12345', 302)

        html, lines_map = docx_parse_document(doc, strict=True)

        self.assertNotIn('data-sticky-until', html)

    def test_heading_bookmarks_with_leading_underscore_are_sticky(self):
        doc = docx.Document()
        p = doc.add_paragraph('Exercise text', style='Intense Quote')
        p._p.append(self._hyperlink('_Упражнение_на_чтение', 'Exercise link'))
        self._add_bookmark(doc.add_paragraph('Next heading'), '_Упражнение_на_чтение', 303)

        html, lines_map = docx_parse_document(doc, strict=True)

        self.assertIn('data-sticky-until="_Упражнение_на_чтение"', html)

    def test_complex_field_hyperlink_makes_image_sticky(self):
        from docx.oxml import parse_xml
        from docx.oxml.ns import nsdecls
        doc = docx.Document()
        p = doc.add_paragraph()
        p.add_run('![')
        p._p.append(parse_xml(f'<w:r {nsdecls("w")}><w:fldChar w:fldCharType="begin"/></w:r>'))
        p._p.append(parse_xml(f'<w:r {nsdecls("w")}><w:instrText xml:space="preserve"> HYPERLINK \\l "_Открытый_звук" </w:instrText></w:r>'))
        p._p.append(parse_xml(f'<w:r {nsdecls("w")}><w:fldChar w:fldCharType="separate"/></w:r>'))
        p.add_run('Photo title')
        p._p.append(parse_xml(f'<w:r {nsdecls("w")}><w:fldChar w:fldCharType="end"/></w:r>'))
        p.add_run('](https://example.com/photo.jpg "Photo alt")')
        self._add_bookmark(doc.add_paragraph('Next heading'), '_Открытый_звук', 401)

        html, lines_map = docx_parse_document(doc, strict=True)
        self.assertIn('<div data-sticky-until="_Открытый_звук">', html)
        self.assertIn('<img title="Photo title" src="https://example.com/photo.jpg" alt="Photo alt">', html)

    def test_html_block_with_leading_style_attaches_sticky_to_container_div(self):
        doc = docx.Document()
        p = doc.add_paragraph('[')
        p._p.append(self._hyperlink('_Анкета', 'HTML'))
        p.add_run(']<style>.foo{color:red;}</style><div class="lf-container"><img src="test.jpg"></div>[/HTML]')
        self._add_bookmark(doc.add_paragraph('Next heading'), '_Анкета', 402)

        html, lines_map = docx_parse_document(doc, strict=True)
        self.assertNotIn('<style data-sticky-until', html)
        self.assertIn('<div class="lf-container" data-sticky-until="_Анкета">', html)

    def test_pipe_marker_split_with_nested_tags(self):
        from docx.oxml import parse_xml
        from docx.oxml.ns import nsdecls
        doc = docx.Document()
        p = doc.add_paragraph()
        p.add_run('— Première ')
        bold_run = p.add_run('partie || seconde')
        bold_run.bold = True
        p.add_run(' partie.')
        p.add_comment('Teacher 0 1000', rangeStart=0, rangeEnd=2)
        bold_run._r.addprevious(parse_xml(f'<w:bookmarkStart {nsdecls("w")} w:id="303" w:name="signet"/>'))
        bold_run._r.addnext(parse_xml(f'<w:bookmarkEnd {nsdecls("w")} w:id="303"/>'))

        html, lines_map = docx_parse_document(doc)

        self.assertEqual(len(lines_map), 2)
        first, second = lines_map
        self.assertNotIn('||', html)
        self.assertEqual(html.count('id="signet"'), 1)
        self.assertIn(
            f'— Première <span id="signet"><strong>partie</strong></span></span> '
            f'<span class="transcript-line" id="{second["id"]}" data-start="{second["start"]}" '
            f'data-end="{second["end"]}" data-speaker="Teacher"><span><strong>seconde</strong></span>',
            html
        )
        self.assertIn(f'id="{first["id"]}" data-start="0" data-end="{first["end"]}"', html)

    def test_raw_html_block_is_not_altered_by_text_processing(self):
        doc = docx.Document()
        doc.add_paragraph('[HTML]<div class="btn -primary"><script>if (a || b) { x = 1 -y; }</script></div>[/HTML]')
        doc.add_paragraph('Texte || normal ?')

        html, lines_map = docx_parse_document(doc)

        self.assertIn('<div class="btn -primary"><script>if (a || b) { x = 1 -y; }</script></div>', html)
        # Markers outside transcript lines are kept, typography is still applied to text
        self.assertIn('<p>Texte || normal&nbsp;?</p>', html)

    def test_json_for_script(self):
        import json
        from home.utils import json_for_script
        value = [{'bookmark': None, 'is_split': True, 'text': '</script><b>&'}]
        result = json_for_script(value)
        self.assertNotIn('<', result)
        self.assertNotIn('&', result)
        self.assertIn('null', result)
        self.assertIn('true', result)
        self.assertEqual(json.loads(result), value)

    def test_sticky_anchor_reverse_order_logs_warning(self):
        from docx.oxml import parse_xml
        from docx.oxml.ns import nsdecls
        import logging
        doc = docx.Document()

        # Paragraph 1: target bookmark is placed here (BEFORE the callout)
        p1 = doc.add_paragraph('Paragraphe initial avec signet.')
        b_start = parse_xml(f'<w:bookmarkStart {nsdecls("w")} w:id="201" w:name="upstream_anchor"/>')
        b_end = parse_xml(f'<w:bookmarkEnd {nsdecls("w")} w:id="201"/>')
        p1._p.insert(0, b_start)
        p1._p.append(b_end)

        # Paragraph 2: Callout referencing upstream_anchor
        p_callout = doc.add_paragraph(style='Intense Quote')
        hl_xml = parse_xml(
            f'<w:hyperlink {nsdecls("w")} w:anchor="upstream_anchor">'
            f'  <w:r><w:t>Règle pointant vers le haut</w:t></w:r>'
            f'</w:hyperlink>'
        )
        p_callout._p.append(hl_xml)

        with self.assertLogs('home.utils', level='WARNING') as cm:
            html, lines_map = docx_parse_document(doc)

        self.assertTrue(any('upstream_anchor' in msg for msg in cm.output))




    def test_overlapping_timing_and_editorial_comments_nest_correctly(self):
        import lxml.html
        doc = docx.Document()
        p = doc.add_paragraph()
        p.add_run('— Alpha ')
        p.add_run('beta ')
        p.add_run('gamma')
        # Timing range covers runs 0-1, editorial range covers runs 1-2 (overlap on "beta ")
        self._comment_runs(p, 'Teacher 0 1000', 0, 1)
        self._comment_runs(p, 'Note', 1, 2, author='editor')

        html, lines_map = docx_parse_document(doc)

        tree = lxml.html.fragment_fromstring(html, create_parent='div')
        lines = tree.find_class('transcript-line')
        self.assertEqual(len(lines), 1)
        self.assertEqual(lines[0].text_content(), '— Alpha beta ')
        notes = tree.find_class('arnaud-comment')
        self.assertEqual(''.join(n.text_content() for n in notes), 'beta gamma')
        self.assertTrue(all(n.get('data-content') == 'Note' for n in notes))

    def test_html_block_keeps_inline_breaks_and_strips_structural_ones(self):
        doc = docx.Document()
        p = doc.add_paragraph()
        p.add_run('[HTML]<table>').add_break()
        p.add_run('<tr><td><b>Je suis</b>').add_break()
        p.add_run('tu es</td></tr></table><div class="lf-container">Ligne 1')
        run = p.add_run()
        run.add_break()
        run.add_break()
        p.add_run('Ligne 2</div>[/HTML]')

        html, lines_map = docx_parse_document(doc)

        self.assertIn('<table><tr><td>', html)
        self.assertIn('<b>Je suis</b><br/>tu es', html)
        self.assertIn('Ligne 1<br/><br/>Ligne 2', html)

    def test_editorial_comment_inside_html_block_keeps_escaping_and_sticky_container(self):
        doc = docx.Document()
        p = doc.add_paragraph('[')
        p._p.append(self._hyperlink('_Tableau', 'HTML'))
        p.add_run('] ')
        commented = p.add_run('<div class="lf-container"><p>Bonjour</p></div>')
        commented.add_comment('a "b" & c', author='editor')
        p.add_run('[/HTML]')
        self._add_bookmark(doc.add_paragraph('Next heading'), '_Tableau', 403)

        html, lines_map = docx_parse_document(doc, strict=True)

        self.assertIn('data-content="a &quot;b&quot; &amp; c"', html)
        self.assertIn('<div class="lf-container" data-sticky-until="_Tableau">', html)
        self.assertNotIn('arnaud-comment" data-container="body" data-content="a &quot;b&quot; &amp; c" '
                         'data-placement="top" data-toggle="popover" data-sticky-until', html)

    def test_editorial_comment_ending_inside_timing_range_is_closed_after_it(self):
        import lxml.html
        doc = docx.Document()
        p = doc.add_paragraph()
        p.add_run('Remarque ')
        p.add_run('— Alpha ')
        p.add_run('beta')
        # Editorial range covers runs 0-1, timing range covers runs 1-2 (overlap on "— Alpha ")
        self._comment_runs(p, 'Note', 0, 1, author='editor')
        self._comment_runs(p, 'Teacher 0 1000', 1, 2)

        html, lines_map = docx_parse_document(doc)

        self.assertEqual(len(lines_map), 1)
        self.assertEqual(html.count(f'id="{lines_map[0]["id"]}"'), 1)
        tree = lxml.html.fragment_fromstring(html, create_parent='div')
        lines = tree.find_class('transcript-line')
        self.assertEqual(len(lines), 1)
        self.assertEqual(lines[0].text_content(), '— Alpha beta')
        notes = tree.find_class('arnaud-comment')
        self.assertTrue(notes[0].text_content().startswith('Remarque — Alpha '))


class ConvertTranscriptCommandTestCase(TestCase):
    def setUp(self):
        import tempfile
        self._tmp = tempfile.TemporaryDirectory()
        self.docx_dir = os.path.join(self._tmp.name, 'DOCX')
        self.output_dir = os.path.join(self._tmp.name, 'JSON')
        os.makedirs(self.docx_dir)

    def tearDown(self):
        self._tmp.cleanup()

    def _write_docx(self, name):
        doc = docx.Document()
        p = doc.add_paragraph('— Bonjour !')
        p.runs[0].add_comment('Teacher 0 1000')
        path = os.path.join(self.docx_dir, name)
        doc.save(path)
        return path

    def test_extract_lesson_number_ignores_dates(self):
        from home.management.commands.convert_transcript import extract_lesson_number
        self.assertEqual(extract_lesson_number('2020-11-19 - Lecon_003_Xv.docx'), 3)
        self.assertEqual(extract_lesson_number('Lecon_004_V.docx'), 4)
        self.assertEqual(extract_lesson_number('French-012-admin.docx'), 12)
        self.assertEqual(extract_lesson_number('Lesson 7.docx'), 7)
        self.assertEqual(extract_lesson_number('Урок_3.docx'), 3)
        self.assertEqual(extract_lesson_number('Leçon_15.docx'), 15)
        self.assertIsNone(extract_lesson_number('notes 2020-11-19.docx'))

    def test_resolve_docx_path_by_number_does_not_match_dates(self):
        from home.management.commands.convert_transcript import resolve_docx_path
        dated = self._write_docx('2003-01-01 - Lecon_007.docx')
        wanted = self._write_docx('2020-11-19 - Lecon_003_Xv.docx')
        self.assertEqual(resolve_docx_path('3', self.docx_dir), wanted)
        self.assertEqual(resolve_docx_path('7', self.docx_dir), dated)
        self.assertEqual(resolve_docx_path('2003-01-01 - Lecon_007', self.docx_dir), dated)
        self.assertIsNone(resolve_docx_path('42', self.docx_dir))

    def test_update_lesson_page_stores_list_and_enables_transcript(self):
        from types import SimpleNamespace
        from unittest import mock
        from home.management.commands.convert_transcript import update_lesson_page
        from home.utils import TranscriptList
        page = SimpleNamespace(transcript='[]', transcript_html='', has_transcript=False, save=mock.Mock())
        lines_map = TranscriptList([{'id': 'line-1', 'start': 0, 'end': 1000}])

        update_lesson_page(page, '<p>html</p>', lines_map)

        self.assertEqual(page.transcript, [{'id': 'line-1', 'start': 0, 'end': 1000}])
        self.assertIs(type(page.transcript), list)
        self.assertEqual(page.transcript_html, '<p>html</p>')
        self.assertTrue(page.has_transcript)
        page.save.assert_called_once_with()

    def test_command_writes_admin_and_masterfile_outputs(self):
        import json
        from django.core.management import call_command
        self._write_docx('2020-11-19 - Lecon_005.docx')

        call_command('convert_transcript', '5', docx_dir=self.docx_dir, output_dir=self.output_dir,
                     stdout=io.StringIO())

        prefix = os.path.join(self.output_dir, 'French-005')
        with open(prefix + '-admin.html', encoding='utf-8') as f:
            self.assertIn('transcript-line', f.read())
        with open(prefix + '-admin.json', encoding='utf-8') as f:
            self.assertEqual(len(json.load(f)), 1)
        with open(prefix + '-masterfile.json', encoding='utf-8') as f:
            self.assertEqual(json.load(f)['lesson_number'], 5)

    def test_command_reports_missing_file(self):
        from django.core.management import call_command, CommandError
        with self.assertRaises(CommandError):
            call_command('convert_transcript', '99', docx_dir=self.docx_dir, output_dir=self.output_dir,
                         stdout=io.StringIO())
